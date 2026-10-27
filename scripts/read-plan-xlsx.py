"""Isolated XLSX-to-CSV feasibility reader; NOT part of the deployed JS application.
No writes, macros, formula evaluation, external links, or source mutation.
Caller supplies the exact existing academic/club CSV header, and uses existing
server preview/confirmation for all business validation. Not a malware scanner.
"""
import csv
import datetime as dt
import io
import math
import zipfile
from pathlib import PurePosixPath
from openpyxl import load_workbook

MAX_INPUT = 1_000_000
MAX_EXPANDED = 4_000_000
MAX_TEXT = 30_000

def read_plan_xlsx(content, columns, sheet=None):
    if not isinstance(content, bytes) or not content or len(content) > MAX_INPUT:
        raise ValueError('Choose an .xlsx workbook no larger than 1 MB.')
    if not isinstance(columns, list) or not columns or len(columns) > 40 or any(not isinstance(c, str) or not c for c in columns) or len(set(columns)) != len(columns):
        raise ValueError('Supply the exact unique existing import columns.')
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as error:
        raise ValueError('Not a readable .xlsx workbook; legacy or encrypted files are unsupported.') from error
    with archive:
        entries = archive.infolist()
        names = [entry.filename for entry in entries]
        if len(entries) > 200 or len(names) != len(set(names)) or sum(e.file_size for e in entries) > MAX_EXPANDED:
            raise ValueError('Workbook archive exceeds the supported limits or has duplicate entries.')
        for entry in entries:
            parts = PurePosixPath(entry.filename).parts
            if entry.flag_bits & 1 or entry.filename.startswith('/') or '..' in parts or '\\' in entry.filename:
                raise ValueError('Encrypted or unsafe workbook archive.')
            if entry.file_size > 2_000_000 or entry.file_size > max(1, entry.compress_size) * 200:
                raise ValueError('Workbook expansion exceeds supported limits.')
            lower = entry.filename.lower()
            if any(k in lower for k in ['vbaproject', 'externallinks/', 'embeddings/', 'activex/']):
                raise ValueError('Macros, external links and embedded objects are unsupported.')
            if lower.endswith(('.xml', '.rels')):
                xml = archive.read(entry)
                if b'<!DOCTYPE' in xml.upper() or b'<!ENTITY' in xml.upper() or b'macroenabled' in xml.lower():
                    raise ValueError('Unsupported workbook XML or macro content.')
        if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
            raise ValueError('Not an .xlsx workbook.')
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=False, keep_links=False)
    except Exception as error:
        raise ValueError('The workbook could not be read.') from error
    try:
        if sheet is None:
            if len(workbook.worksheets) != 1:
                raise ValueError('Choose one worksheet explicitly; multiple sheets are not combined.')
            worksheet = workbook.worksheets[0]
        else:
            if sheet not in workbook.sheetnames:
                raise ValueError('Selected worksheet does not exist.')
            worksheet = workbook[sheet]
        if worksheet.sheet_state != 'visible' or worksheet.merged_cells.ranges:
            raise ValueError('Selected worksheet must be visible and have no merged cells.')
        if worksheet.max_row > 51 or worksheet.max_column != len(columns):
            raise ValueError('Use at most 50 data rows and the exact import columns; remove excess formatting.')
        if any(d.hidden for d in worksheet.row_dimensions.values()) or any(d.hidden for d in worksheet.column_dimensions.values()):
            raise ValueError('Hidden rows or columns are unsupported in the selected worksheet.')
        header = [cell.value for cell in worksheet[1]]
        if header != columns:
            raise ValueError('Worksheet header must match the existing import header exactly.')
        rows = [columns]
        for row in worksheet.iter_rows(min_row=2, max_col=len(columns)):
            if all(cell.value is None for cell in row):
                continue
            values = []
            for column, cell in zip(columns, row):
                value = cell.value
                if cell.data_type in ('f', 'e') or cell.hyperlink:
                    raise ValueError('Formulas, spreadsheet errors and hyperlinks are unsupported.')
                if value is None:
                    values.append('')
                elif isinstance(value, bool):
                    raise ValueError('Boolean cells are unsupported; use explicit text.')
                elif isinstance(value, (dt.datetime, dt.date)):
                    if column not in ('start', 'end', 'reportDue') or not 1970 <= value.year <= 2100:
                        raise ValueError('Use supported calendar dates only in date columns.')
                    if isinstance(value, dt.datetime) and value.time() != dt.time():
                        raise ValueError('Date columns cannot contain a time of day.')
                    values.append(value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat())
                elif isinstance(value, str):
                    if len(value) > 2000:
                        raise ValueError('Cell exceeds 2,000 characters.')
                    values.append(value)
                elif isinstance(value, (int, float)) and math.isfinite(value):
                    if column in ('externalId', 'lead'):
                        raise ValueError('Identifiers must be text cells to preserve leading zeros.')
                    values.append(str(int(value)) if float(value).is_integer() else str(value))
                else:
                    raise ValueError('Unsupported cell value.')
            rows.append(values)
        if len(rows) < 2:
            raise ValueError('Include at least one data row.')
        output = io.StringIO(newline='')
        csv.writer(output, lineterminator='\n').writerows(rows)
        result = output.getvalue()
        if len(result) > MAX_TEXT:
            raise ValueError('Converted CSV exceeds 30,000 characters.')
        return {'sheet': worksheet.title, 'rows': len(rows)-1, 'csv': result}
    finally:
        workbook.close()

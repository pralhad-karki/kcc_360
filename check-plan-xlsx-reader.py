import sys
sys.dont_write_bytecode=True
import csv
import datetime as dt
import importlib.util
import io
import unittest
import zipfile
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
spec=importlib.util.spec_from_file_location('reader','scripts/read-plan-xlsx.py');reader=importlib.util.module_from_spec(spec);spec.loader.exec_module(reader)
columns=['externalId','title','start','end','reportDue','lead','budget']
def workbook():
    w=Workbook();w.active.title='Activities';w.active.append(columns);w.active.append(['ROW-001','Workshop, "community"\nreview',dt.date(2026,11,1),'2026-11-02','2026-11-03','s1',125.5]);return w

def packed(w):
    output=io.BytesIO();w.save(output);return output.getvalue()

class ReaderChecks(unittest.TestCase):
    def test_roundtrip_epochs_and_preservation(self):
        for epoch in [None,CALENDAR_MAC_1904]:
            w=workbook()
            if epoch:w.epoch=epoch
            raw=packed(w);result=reader.read_plan_xlsx(raw,columns);rows=list(csv.reader(io.StringIO(result['csv'])));self.assertEqual(rows[0],columns);self.assertEqual(rows[1][0],'ROW-001');self.assertEqual(rows[1][1],w.active['B2'].value);self.assertEqual(rows[1][2],'2026-11-01');self.assertEqual(rows[1][-1],'125.5');self.assertEqual(result['rows'],1);self.assertEqual(raw,packed(w))
    def test_bad_cells_and_layout(self):
        mutations=[lambda w:setattr(w.active['A2'],'value',1),lambda w:setattr(w.active['B2'],'value','=1+1'),lambda w:setattr(w.active['G2'],'value',True),lambda w:setattr(w.active['B2'],'value','x'*2001),lambda w:setattr(w.active['C2'],'value',dt.datetime(2026,11,1,12)),lambda w:setattr(w.active['B2'],'value',dt.date(2026,11,1)),lambda w:setattr(w.active['A1'],'value','wrong'),lambda w:w.active.merge_cells('A2:B2'),lambda w:setattr(w.active.row_dimensions[2],'hidden',True),lambda w:setattr(w.active.column_dimensions['B'],'hidden',True),lambda w:setattr(w.active['B2'],'hyperlink','https://example.invalid'),lambda w:setattr(w.active['H2'],'value','extra')]
        for mutate in mutations:
            with self.subTest(mutate=mutate):
                w=workbook();mutate(w)
                with self.assertRaises(ValueError):reader.read_plan_xlsx(packed(w),columns)
    def test_selection_and_limits(self):
        w=workbook();w.create_sheet('Other');raw=packed(w)
        with self.assertRaisesRegex(ValueError,'explicitly'):reader.read_plan_xlsx(raw,columns)
        self.assertEqual(reader.read_plan_xlsx(raw,columns,'Activities')['rows'],1)
        with self.assertRaises(ValueError):reader.read_plan_xlsx(raw,columns,'Missing')
        w.active.sheet_state='hidden'
        with self.assertRaises(ValueError):reader.read_plan_xlsx(packed(w),columns,'Activities')
        w=workbook()
        for i in range(50):w.active.append(['ROW-'+str(i),'Text','2026-11-01','2026-11-02','2026-11-03','s1',0])
        with self.assertRaises(ValueError):reader.read_plan_xlsx(packed(w),columns)
        with self.assertRaises(ValueError):reader.read_plan_xlsx(b'not excel',columns)
        with self.assertRaises(ValueError):reader.read_plan_xlsx(b'x'*1000001,columns)
    def test_unsafe_archive(self):
        raw=packed(workbook())
        for path,value in [('xl/vbaProject.bin',b'code'),('xl/externalLinks/test.xml',b'x'),('../escape',b'x'),('evil.xml',b'<!DOCTYPE x>')]:
            output=io.BytesIO()
            with zipfile.ZipFile(io.BytesIO(raw)) as original,zipfile.ZipFile(output,'w') as target:
                for item in original.infolist():target.writestr(item,original.read(item.filename))
                target.writestr(path,value)
            with self.assertRaises(ValueError):reader.read_plan_xlsx(output.getvalue(),columns)

if __name__=='__main__':unittest.main()
